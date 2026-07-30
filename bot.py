import asyncio
import hashlib
import html
import io
import json
import logging
import os
import re
import sqlite3
import time
import uuid
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

import aiohttp
from dotenv import load_dotenv

try:
    import openpyxl
except ImportError:
    openpyxl = None

from aiogram import BaseMiddleware, Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.filters import Command, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    CallbackQuery,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    Message,
    TelegramObject,
)


# ============================================================
# CONFIG
# ============================================================

load_dotenv()

BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()

SMS_API_DOMAIN = os.getenv("SMS_API_DOMAIN", "").strip().rstrip("/")
SMS_API_USERNAME = os.getenv("SMS_API_USERNAME", "").strip()
SMS_API_KEY = os.getenv("SMS_API_KEY", "").strip()
SMS_API_SIGN_TYPE = os.getenv("SMS_API_SIGN_TYPE", "MD5").strip().upper()
SMS_API_SP_NUMBER = os.getenv("SMS_API_SP_NUMBER", "").strip()

SMS_API_SUBMITTAL_URL = os.getenv("SMS_API_SUBMITTAL_URL", "").strip()
SMS_API_RECORDS_URL = os.getenv("SMS_API_RECORDS_URL", "").strip()
SMS_API_BALANCE_URL = os.getenv("SMS_API_BALANCE_URL", "").strip()

if SMS_API_DOMAIN:
    if not SMS_API_SUBMITTAL_URL:
        SMS_API_SUBMITTAL_URL = f"{SMS_API_DOMAIN}/ta-sms/openapi/submittal"
    if not SMS_API_RECORDS_URL:
        SMS_API_RECORDS_URL = f"{SMS_API_DOMAIN}/ta-sms/openapi/records"
    if not SMS_API_BALANCE_URL:
        SMS_API_BALANCE_URL = f"{SMS_API_DOMAIN}/ta-sms/openapi/balance"

BATCH_SIZE = int(os.getenv("BATCH_SIZE", "100"))
MAX_NUMBERS = int(os.getenv("MAX_NUMBERS", "5000"))
REQUEST_TIMEOUT = int(os.getenv("REQUEST_TIMEOUT", "30"))
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "20"))

SMS_API_KEEP_PLUS = os.getenv("SMS_API_KEEP_PLUS", "false").lower() in ("1", "true", "yes", "on")
SMS_API_TIMESTAMP_UNIT = os.getenv("SMS_API_TIMESTAMP_UNIT", "ms").lower().strip()
SMS_API_SIGN_UPPER = os.getenv("SMS_API_SIGN_UPPER", "true").lower() in ("1", "true", "yes", "on")

LOG_TIMESTAMPS = os.getenv("LOG_TIMESTAMPS", "false").lower() in ("1", "true", "yes", "on")


def parse_allowed_users(raw: str) -> set[int]:
    raw = raw.strip()
    if not raw:
        return set()

    ids = set()
    for part in re.split(r"[,\s;]+", raw):
        part = part.strip()
        if not part:
            continue
        try:
            ids.add(int(part))
        except ValueError:
            pass
    return ids


ALLOWED_USER_IDS = parse_allowed_users(os.getenv("ALLOWED_USER_IDS", ""))
ADMIN_IDS = parse_allowed_users(os.getenv("ADMIN_IDS", ""))

# Compatibilité ancienne config :
# si ADMIN_IDS est vide mais ALLOWED_USER_IDS existe, ALLOWED_USER_IDS devient admin.
if not ADMIN_IDS and ALLOWED_USER_IDS:
    ADMIN_IDS = set(ALLOWED_USER_IDS)

DB_PATH = os.getenv("DB_PATH", "bot.db").strip() or "bot.db"

SMS_PRICE_SIM = Decimal(os.getenv("SMS_PRICE_SIM", "0.040"))
SMS_PRICE_SENDER = Decimal(os.getenv("SMS_PRICE_SENDER", "0.070"))

if not BOT_TOKEN:
    raise RuntimeError("BOT_TOKEN manquant dans le fichier .env")

if BATCH_SIZE <= 0:
    BATCH_SIZE = 100

if MAX_NUMBERS <= 0:
    MAX_NUMBERS = 5000

if MAX_UPLOAD_MB <= 0:
    MAX_UPLOAD_MB = 20

if SMS_API_TIMESTAMP_UNIT not in ("ms", "s", "sec", "second", "seconds"):
    SMS_API_TIMESTAMP_UNIT = "ms"


# ============================================================
# LOGS
# ============================================================

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
)

logger = logging.getLogger("sms-bot")


# ============================================================
# LICENSES / INTERNAL BALANCE DB
# ============================================================

def now_ts() -> int:
    return int(time.time())


def money(value: Any) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.001"))


def money_fmt(value: Any) -> str:
    return f"{money(value):.3f}"


def db_conn() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_app_db():
    with db_conn() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS users (
                user_id INTEGER PRIMARY KEY,
                balance TEXT NOT NULL DEFAULT '0.000',
                created_at INTEGER NOT NULL
            )
        """)

        conn.execute("""
            CREATE TABLE IF NOT EXISTS licenses (
                user_id INTEGER PRIMARY KEY,
                expires_at INTEGER NOT NULL,
                created_at INTEGER NOT NULL
            )
        """)

        conn.execute("""
            CREATE TABLE IF NOT EXISTS sms_ledger (
                id TEXT PRIMARY KEY,
                campaign_id TEXT NOT NULL,
                user_id INTEGER NOT NULL,
                route TEXT NOT NULL,
                phone TEXT NOT NULL,
                amount TEXT NOT NULL,
                status TEXT NOT NULL,
                provider_status TEXT,
                provider_message_id TEXT,
                created_at INTEGER NOT NULL,
                updated_at INTEGER NOT NULL
            )
        """)

        conn.commit()


def is_admin(user_id: int) -> bool:
    return user_id in ADMIN_IDS


def ensure_user(user_id: int):
    with db_conn() as conn:
        conn.execute(
            """
            INSERT OR IGNORE INTO users(user_id, balance, created_at)
            VALUES (?, ?, ?)
            """,
            (user_id, "0.000", now_ts()),
        )
        conn.commit()


def delete_expired_licenses() -> int:
    current = now_ts()

    with db_conn() as conn:
        cur = conn.execute(
            "DELETE FROM licenses WHERE expires_at <= ?",
            (current,),
        )
        conn.commit()
        return cur.rowcount or 0


def has_valid_license(user_id: int) -> bool:
    if is_admin(user_id):
        return True

    delete_expired_licenses()

    with db_conn() as conn:
        row = conn.execute(
            "SELECT expires_at FROM licenses WHERE user_id = ?",
            (user_id,),
        ).fetchone()

    if not row:
        return False

    return int(row["expires_at"]) > now_ts()


def parse_duration_to_seconds(raw: str) -> int:
    """
    Formats acceptés :
    30      = 30 jours
    30d     = 30 jours
    12h     = 12 heures
    60m     = 60 minutes
    2w      = 2 semaines
    """
    raw = raw.strip().lower()

    match = re.fullmatch(r"(\d+)([mhdw]?)", raw)
    if not match:
        raise ValueError("Durée invalide. Exemple : 30d, 12h, 60m, 2w")

    amount = int(match.group(1))
    unit = match.group(2) or "d"

    if amount <= 0:
        raise ValueError("La durée doit être supérieure à 0.")

    if unit == "m":
        return amount * 60
    if unit == "h":
        return amount * 3600
    if unit == "d":
        return amount * 86400
    if unit == "w":
        return amount * 7 * 86400

    raise ValueError("Unité invalide.")


def grant_license(user_id: int, duration_raw: str) -> int:
    ensure_user(user_id)

    seconds = parse_duration_to_seconds(duration_raw)
    expires_at = now_ts() + seconds

    with db_conn() as conn:
        conn.execute(
            """
            INSERT INTO licenses(user_id, expires_at, created_at)
            VALUES (?, ?, ?)
            ON CONFLICT(user_id)
            DO UPDATE SET expires_at = excluded.expires_at
            """,
            (user_id, expires_at, now_ts()),
        )
        conn.commit()

    return expires_at


def revoke_license(user_id: int):
    with db_conn() as conn:
        conn.execute(
            "DELETE FROM licenses WHERE user_id = ?",
            (user_id,),
        )
        conn.commit()


def get_license_expiry(user_id: int) -> Optional[int]:
    delete_expired_licenses()

    with db_conn() as conn:
        row = conn.execute(
            "SELECT expires_at FROM licenses WHERE user_id = ?",
            (user_id,),
        ).fetchone()

    if not row:
        return None

    return int(row["expires_at"])


def get_user_balance(user_id: int) -> Decimal:
    ensure_user(user_id)

    with db_conn() as conn:
        row = conn.execute(
            "SELECT balance FROM users WHERE user_id = ?",
            (user_id,),
        ).fetchone()

    return money(row["balance"])


def set_user_balance(user_id: int, amount: Decimal) -> Decimal:
    ensure_user(user_id)

    amount = money(amount)

    if amount < money("0"):
        raise ValueError("Le solde ne peut pas être négatif.")

    with db_conn() as conn:
        conn.execute(
            "UPDATE users SET balance = ? WHERE user_id = ?",
            (str(amount), user_id),
        )
        conn.commit()

    return amount


def add_user_balance(user_id: int, amount: Decimal) -> Decimal:
    ensure_user(user_id)

    if amount < money("0"):
        raise ValueError("Le montant doit être positif.")

    current = get_user_balance(user_id)
    new_balance = money(current + money(amount))

    set_user_balance(user_id, new_balance)

    return new_balance


def remove_user_balance(user_id: int, amount: Decimal) -> Decimal:
    ensure_user(user_id)

    if amount < money("0"):
        raise ValueError("Le montant doit être positif.")

    current = get_user_balance(user_id)

    if current < amount:
        raise ValueError(
            f"Solde insuffisant. Disponible: ${money_fmt(current)}, retrait demandé: ${money_fmt(amount)}"
        )

    new_balance = money(current - amount)

    set_user_balance(user_id, new_balance)

    return new_balance


def get_sms_price(route: str) -> Decimal:
    if route == "sim":
        return money(SMS_PRICE_SIM)

    if route == "sender":
        return money(SMS_PRICE_SENDER)

    raise ValueError(f"Route inconnue : {route}")


def get_sms_total_cost(route: str, count: int) -> Decimal:
    return money(get_sms_price(route) * count)


def reserve_sms_balance(
    user_id: int,
    route: str,
    phones: List[str],
    campaign_id: str,
) -> Tuple[List[str], Decimal]:
    """
    Débite avant envoi et crée un ledger par SMS.
    Si API refuse ou statut final failed => refund.
    """
    ensure_user(user_id)

    price = get_sms_price(route)
    total_cost = get_sms_total_cost(route, len(phones))
    current_balance = get_user_balance(user_id)

    if current_balance < total_cost:
        raise ValueError(
            f"Solde insuffisant. Requis: ${money_fmt(total_cost)}, disponible: ${money_fmt(current_balance)}"
        )

    created = now_ts()
    ledger_ids: List[str] = []

    with db_conn() as conn:
        new_balance = money(current_balance - total_cost)

        conn.execute(
            "UPDATE users SET balance = ? WHERE user_id = ?",
            (str(new_balance), user_id),
        )

        for phone in phones:
            ledger_id = uuid.uuid4().hex
            ledger_ids.append(ledger_id)

            conn.execute(
                """
                INSERT INTO sms_ledger(
                    id, campaign_id, user_id, route, phone, amount,
                    status, provider_status, provider_message_id,
                    created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    ledger_id,
                    campaign_id,
                    user_id,
                    route,
                    phone,
                    str(price),
                    "reserved",
                    None,
                    None,
                    created,
                    created,
                ),
            )

        conn.commit()

    return ledger_ids, total_cost


def update_ledger_status(
    ledger_ids: List[str],
    status: str,
    provider_status: Optional[str] = None,
    provider_message_id: Optional[str] = None,
):
    if not ledger_ids:
        return

    updated = now_ts()
    placeholders = ",".join(["?"] * len(ledger_ids))

    with db_conn() as conn:
        conn.execute(
            f"""
            UPDATE sms_ledger
            SET status = ?,
                provider_status = COALESCE(?, provider_status),
                provider_message_id = COALESCE(?, provider_message_id),
                updated_at = ?
            WHERE id IN ({placeholders})
            """,
            [status, provider_status, provider_message_id, updated, *ledger_ids],
        )
        conn.commit()


def refund_ledger_ids(
    ledger_ids: List[str],
    provider_status: Optional[str] = None,
) -> Tuple[int, Decimal]:
    """
    Refund uniquement les SMS pas encore delivered/refunded.
    """
    if not ledger_ids:
        return 0, money("0")

    placeholders = ",".join(["?"] * len(ledger_ids))

    with db_conn() as conn:
        rows = conn.execute(
            f"""
            SELECT id, user_id, amount, status
            FROM sms_ledger
            WHERE id IN ({placeholders})
            AND status IN ('reserved', 'accepted')
            """,
            ledger_ids,
        ).fetchall()

        if not rows:
            return 0, money("0")

        refund_by_user: Dict[int, Decimal] = {}
        refundable_ids: List[str] = []

        for row in rows:
            uid = int(row["user_id"])
            amount = money(row["amount"])
            refundable_ids.append(row["id"])
            refund_by_user[uid] = refund_by_user.get(uid, money("0")) + amount

        for uid, amount in refund_by_user.items():
            user_row = conn.execute(
                "SELECT balance FROM users WHERE user_id = ?",
                (uid,),
            ).fetchone()

            current_balance = money(user_row["balance"]) if user_row else money("0")
            new_balance = money(current_balance + amount)

            conn.execute(
                "UPDATE users SET balance = ? WHERE user_id = ?",
                (str(new_balance), uid),
            )

        refund_placeholders = ",".join(["?"] * len(refundable_ids))

        conn.execute(
            f"""
            UPDATE sms_ledger
            SET status = 'refunded',
                provider_status = COALESCE(?, provider_status),
                updated_at = ?
            WHERE id IN ({refund_placeholders})
            """,
            [provider_status, now_ts(), *refundable_ids],
        )

        conn.commit()

    total_refunded = sum((money(row["amount"]) for row in rows), money("0"))
    return len(rows), money(total_refunded)


def mark_campaign_batch_accepted(ledger_ids: List[str]):
    update_ledger_status(ledger_ids, "accepted", "accepted_api")


def is_delivered_status(status: Any) -> bool:
    s = str(status).strip().lower()

    return s in {
        "2",
        "delivered",
        "deliver",
        "delivrd",
        "success",
        "successful",
        "ok",
        "sent_success",
        "delivery_success",
    }


def is_failed_status(status: Any) -> bool:
    s = str(status).strip().lower()

    return s in {
        "3",
        "4",
        "5",
        "failed",
        "fail",
        "rejected",
        "reject",
        "undeliv",
        "undelivered",
        "expired",
        "error",
        "delivery_failed",
        "not_delivered",
        "blacklist",
    }


def settle_one_record_by_phone(phone: str, provider_status: str) -> Optional[str]:
    """
    Match simple via numéro.
    Delivered => marque delivered.
    Failed => refund.
    """
    normalized = normalize_number(str(phone))

    if not normalized:
        return None

    with db_conn() as conn:
        row = conn.execute(
            """
            SELECT id
            FROM sms_ledger
            WHERE phone = ?
            AND status IN ('reserved', 'accepted')
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (normalized,),
        ).fetchone()

    if not row:
        return None

    ledger_id = row["id"]

    if is_delivered_status(provider_status):
        update_ledger_status([ledger_id], "delivered", provider_status)
        return "delivered"

    if is_failed_status(provider_status):
        refund_ledger_ids([ledger_id], provider_status)
        return "refunded"

    return "pending"


def extract_record_phone_status_pairs(data: Any) -> List[Tuple[str, str]]:
    """
    Extraction générique des records API.
    Cherche récursivement des objets contenant :
    phone/msisdn/mobile/number/to + status/state/deliveryStatus...
    """
    pairs: List[Tuple[str, str]] = []

    phone_keys = {
        "phone",
        "phones",
        "mobile",
        "msisdn",
        "number",
        "to",
        "recipient",
    }

    status_keys = {
        "status",
        "state",
        "smsStatus",
        "deliveryStatus",
        "deliverStatus",
        "reportStatus",
        "sendStatus",
        "code",
    }

    if isinstance(data, dict):
        phone_value = None
        status_value = None

        for key, value in data.items():
            if key in phone_keys and isinstance(value, (str, int, float)):
                phone_value = str(value)

            if key in status_keys and isinstance(value, (str, int, float)):
                status_value = str(value)

        if phone_value and status_value:
            pairs.append((phone_value, status_value))

        for value in data.values():
            pairs.extend(extract_record_phone_status_pairs(value))

    elif isinstance(data, list):
        for item in data:
            pairs.extend(extract_record_phone_status_pairs(item))

    return pairs


def auto_settle_from_records(data: Any) -> Dict[str, int]:
    pairs = extract_record_phone_status_pairs(data)

    result = {
        "found": len(pairs),
        "delivered": 0,
        "refunded": 0,
        "pending": 0,
    }

    for phone, status in pairs:
        settled = settle_one_record_by_phone(phone, status)

        if settled == "delivered":
            result["delivered"] += 1
        elif settled == "refunded":
            result["refunded"] += 1
        elif settled == "pending":
            result["pending"] += 1

    return result


def get_user_info(user_id: int) -> Dict[str, Any]:
    ensure_user(user_id)

    balance = get_user_balance(user_id)
    expires_at = get_license_expiry(user_id)

    with db_conn() as conn:
        total_sms = conn.execute(
            "SELECT COUNT(*) AS c FROM sms_ledger WHERE user_id = ?",
            (user_id,),
        ).fetchone()["c"]

        delivered_sms = conn.execute(
            "SELECT COUNT(*) AS c FROM sms_ledger WHERE user_id = ? AND status = 'delivered'",
            (user_id,),
        ).fetchone()["c"]

        refunded_sms = conn.execute(
            "SELECT COUNT(*) AS c FROM sms_ledger WHERE user_id = ? AND status = 'refunded'",
            (user_id,),
        ).fetchone()["c"]

        accepted_sms = conn.execute(
            "SELECT COUNT(*) AS c FROM sms_ledger WHERE user_id = ? AND status = 'accepted'",
            (user_id,),
        ).fetchone()["c"]

    return {
        "user_id": user_id,
        "balance": balance,
        "expires_at": expires_at,
        "is_admin": is_admin(user_id),
        "has_license": has_valid_license(user_id),
        "total_sms": total_sms,
        "delivered_sms": delivered_sms,
        "refunded_sms": refunded_sms,
        "accepted_sms": accepted_sms,
    }


async def license_cleanup_loop():
    while True:
        try:
            deleted = delete_expired_licenses()
            if deleted:
                logger.info("Licences expirées supprimées : %s", deleted)
        except Exception:
            logger.exception("Erreur nettoyage licences expirées")

        await asyncio.sleep(3600)


# ============================================================
# BOT / ROUTER
# ============================================================

bot = Bot(
    token=BOT_TOKEN,
    default=DefaultBotProperties(parse_mode=ParseMode.HTML),
)

dp = Dispatcher()
router = Router()
dp.include_router(router)


# ============================================================
# FSM
# ============================================================

class CampaignState(StatesGroup):
    waiting_sender_id = State()
    waiting_numbers = State()
    waiting_text = State()
    waiting_confirmation = State()


# ============================================================
# SECURITY MIDDLEWARE
# ============================================================

class AuthMiddleware(BaseMiddleware):
    async def __call__(self, handler, event: TelegramObject, data: Dict[str, Any]):
        user = data.get("event_from_user")

        if not user:
            return await handler(event, data)

        user_id = user.id
        ensure_user(user_id)
        delete_expired_licenses()

        public_commands = {
            "/start",
            "/campagne",
            "/licence",
            "/license",
            "/solde",
            "/my_balance",
            "/help",
        }

        if isinstance(event, Message):
            text = event.text or ""
            command = text.split(maxsplit=1)[0].lower() if text.startswith("/") else ""

            if command in public_commands:
                return await handler(event, data)

        if is_admin(user_id):
            return await handler(event, data)

        if not has_valid_license(user_id):
            if isinstance(event, Message):
                await event.answer(
                    "⛔ <b>Accès refusé.</b>\n\n"
                    "Tu n'as pas de licence active ou ta licence est expirée.\n"
                    "Contacte l'administrateur."
                )
            elif isinstance(event, CallbackQuery):
                await event.answer(
                    "Licence expirée ou inexistante.",
                    show_alert=True,
                )
            return None

        return await handler(event, data)


router.message.middleware(AuthMiddleware())
router.callback_query.middleware(AuthMiddleware())


# ============================================================
# KEYBOARDS
# ============================================================

def route_keyboard(user_id: int) -> InlineKeyboardMarkup:
    buttons = [
        [InlineKeyboardButton(text="📲 Route SIM - $0.040/SMS", callback_data="route:sim")],
        [InlineKeyboardButton(text="🏷️ Sender ID - $0.070/SMS", callback_data="route:sender")],
        [
            InlineKeyboardButton(text="💰 Mon solde", callback_data="user:balance"),
            InlineKeyboardButton(text="📄 Ma licence", callback_data="user:license"),
        ],
    ]

    if is_admin(user_id):
        buttons.append([
            InlineKeyboardButton(text="👑 Admin", callback_data="admin:panel"),
        ])

    return InlineKeyboardMarkup(inline_keyboard=buttons)


def cancel_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="❌ Annuler", callback_data="campaign:cancel")],
        ],
    )


def confirmation_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="🚀 Envoyer SMS", callback_data="campaign:send")],
            [InlineKeyboardButton(text="✏️ Modifier texte", callback_data="campaign:edit_text")],
            [InlineKeyboardButton(text="🔙 Changer route", callback_data="campaign:restart")],
            [InlineKeyboardButton(text="❌ Annuler", callback_data="campaign:cancel")],
        ],
    )


# ============================================================
# UTILS
# ============================================================

def menu_text() -> str:
    return (
        "📨 <b>Kaplan SMS Bot</b>\n\n"
        "Choisis une route d’envoi :\n\n"
        "📲 <b>Route SIM</b>\n"
        "Prix : <b>$0.040/SMS</b>\n\n"
        "🏷️ <b>Sender ID</b>\n"
        "Prix : <b>$0.070/SMS</b>\n\n"
        "👇 Sélectionne une option :"
    )


def numbers_help_text() -> str:
    return (
        "📞 Envoie maintenant ta liste de numéros.\n\n"
        "Tu peux envoyer :\n"
        "✅ un message direct\n"
        "✅ un fichier <code>.txt</code>\n"
        "✅ un fichier <code>.xlsx</code>\n\n"
        "Formats acceptés :\n"
        "<code>33612345678</code>\n"
        "<code>+33612345678</code>\n"
        "<code>0033612345678</code>\n\n"
        "Les numéros peuvent être séparés par espace, virgule ou retour à la ligne."
    )


def chunk_list(items: List[str], size: int) -> List[List[str]]:
    return [items[i:i + size] for i in range(0, len(items), size)]


def scientific_to_plain_number(value: str) -> Optional[str]:
    value = value.strip().replace(",", ".")

    if not re.fullmatch(r"[+-]?\d+(?:\.\d+)?[eE][+-]?\d+", value):
        return None

    try:
        d = Decimal(value)
        if d == d.to_integral_value():
            return str(d.quantize(Decimal("1")))
        return str(d)
    except (InvalidOperation, ValueError):
        return None


def normalize_number(raw: str) -> Optional[str]:
    if not raw:
        return None

    s = raw.strip()

    sci = scientific_to_plain_number(s)
    if sci:
        s = sci

    s = re.sub(r"[^\d+]", "", s)

    if not s:
        return None

    if s.count("+") > 1:
        return None

    if "+" in s and not s.startswith("+"):
        return None

    if s.startswith("00"):
        if SMS_API_KEEP_PLUS:
            s = "+" + s[2:]
        else:
            s = s[2:]

    elif s.startswith("+"):
        if not SMS_API_KEEP_PLUS:
            s = s[1:]

    digits = s[1:] if s.startswith("+") else s

    if not digits.isdigit():
        return None

    if len(digits) < 8 or len(digits) > 16:
        return None

    return s


def parse_numbers(text: str) -> Tuple[List[str], int, int]:
    parts = re.split(r"[\s,;|]+", text.strip())

    numbers: List[str] = []
    seen = set()
    duplicates = 0
    invalid = 0

    for part in parts:
        part = part.strip()
        if not part:
            continue

        normalized = normalize_number(part)

        if not normalized:
            invalid += 1
            continue

        if normalized in seen:
            duplicates += 1
            continue

        seen.add(normalized)
        numbers.append(normalized)

    return numbers, duplicates, invalid


def cell_value_to_text(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, bool):
        return ""

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, "f")

    value = str(value).strip()

    sci = scientific_to_plain_number(value)
    if sci:
        return sci

    return value


def txt_bytes_to_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue

    return data.decode("utf-8", errors="ignore")


def xlsx_bytes_to_text(data: bytes) -> str:
    if openpyxl is None:
        raise RuntimeError("openpyxl n'est pas installé. Fais : pip install openpyxl")

    buffer = io.BytesIO(data)

    workbook = openpyxl.load_workbook(
        filename=buffer,
        read_only=True,
        data_only=True,
    )

    values: List[str] = []

    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    text = cell_value_to_text(cell)
                    if text:
                        values.append(text)
    finally:
        workbook.close()

    return "\n".join(values)


async def extract_numbers_input(message: Message) -> Tuple[Optional[str], Optional[str], Optional[str]]:
    if message.document:
        document = message.document

        file_name = document.file_name or "fichier"
        safe_file_name = html.escape(file_name)
        ext = os.path.splitext(file_name)[1].lower()

        if ext not in (".txt", ".xlsx"):
            return (
                None,
                None,
                "⚠️ Format non supporté.\n\n"
                "Envoie un fichier <code>.txt</code> ou <code>.xlsx</code>.",
            )

        if document.file_size and document.file_size > MAX_UPLOAD_MB * 1024 * 1024:
            return (
                None,
                None,
                f"⚠️ Fichier trop lourd.\n\nTaille max : <b>{MAX_UPLOAD_MB} MB</b>.",
            )

        buffer = io.BytesIO()
        await bot.download(document, destination=buffer)
        file_data = buffer.getvalue()

        try:
            if ext == ".txt":
                text = txt_bytes_to_text(file_data)
                source = f"📄 TXT : {safe_file_name}"
            else:
                text = xlsx_bytes_to_text(file_data)
                source = f"📊 XLSX : {safe_file_name}"

            if message.caption:
                text += "\n" + message.caption

            return text, source, None

        except Exception as e:
            logger.exception("Erreur lecture fichier numéros")
            return (
                None,
                None,
                f"❌ Impossible de lire le fichier.\n\n"
                f"Erreur : <code>{html.escape(str(e))}</code>",
            )

    if message.text:
        return message.text, "💬 Message direct", None

    return (
        None,
        None,
        "⚠️ Envoie les numéros en message direct, fichier <code>.txt</code> ou fichier <code>.xlsx</code>.",
    )


def make_timestamp(unit: Optional[str] = None) -> str:
    unit = (unit or SMS_API_TIMESTAMP_UNIT).lower().strip()

    if unit in ("s", "sec", "second", "seconds"):
        return str(int(time.time()))

    return str(int(time.time() * 1000))


def truncate_text(text: str, limit: int = 3500) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + "\n... tronqué ..."


def json_pretty(data: Any) -> str:
    try:
        return json.dumps(data, ensure_ascii=False, indent=2)
    except Exception:
        return str(data)


# ============================================================
# SMS API SIGNATURE / REQUESTS
# ============================================================

def sms_api_sign_value(value: Any) -> str:
    if isinstance(value, dict):
        items = []
        for k in sorted(value.keys()):
            items.append(f"{k}={sms_api_sign_value(value[k])}")
        return "{" + ",".join(items) + "}"
    if isinstance(value, list):
        return "[" + ",".join(sms_api_sign_value(v) for v in value) + "]"
    if value is None:
        return ""
    return str(value)


def build_sms_api_signature(params: Dict[str, Any]) -> str:
    sign_params: Dict[str, Any] = {}

    for key, value in params.items():
        if key == "sign":
            continue
        if value is None:
            continue
        if value == "":
            continue
        sign_params[key] = value

    raw = "&".join(
        f"{key}={sms_api_sign_value(sign_params[key])}"
        for key in sorted(sign_params.keys())
    )

    raw = f"{raw}&key={SMS_API_KEY}"
    logger.info("RAW_SIGN=%s", raw)

    sign_type = SMS_API_SIGN_TYPE.upper()

    if sign_type == "MD5":
        digest = hashlib.md5(raw.encode("utf-8")).hexdigest()
    elif sign_type in ("SHA256", "SHA-256"):
        digest = hashlib.sha256(raw.encode("utf-8")).hexdigest()
    else:
        raise RuntimeError(f"signType non supporté : {SMS_API_SIGN_TYPE}")

    if SMS_API_SIGN_UPPER:
        digest = digest.upper()

    return digest


def build_sms_api_headers() -> Dict[str, str]:
    return {
        "Content-Type": "application/json",
        "ta-version": "v2",
    }


def build_base_sms_api_payload(extra: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    payload: Dict[str, Any] = {
        "username": SMS_API_USERNAME,
        "nonceStr": uuid.uuid4().hex,
        "timestamp": make_timestamp(),
        "signType": SMS_API_SIGN_TYPE,
    }

    if extra:
        for key, value in extra.items():
            if value is not None and value != "":
                payload[key] = value

    # IMPORTANT :
    # La signature est calculée APRÈS avoir ajouté content, phones, spNumber, etc.
    payload["sign"] = build_sms_api_signature(payload)
    logger.info("PAYLOAD=%s", payload)

    if LOG_TIMESTAMPS:
        logger.info("SMS_API timestamp=%s", payload["timestamp"])

    return payload


def build_payload(
    numbers: List[str],
    sms_text: str,
    route: str,
    sender_id: Optional[str] = None,
) -> Dict[str, Any]:
    extra: Dict[str, Any] = {
        "content": sms_text,
        "phones": [{"phone": number} for number in numbers],
    }

    if route == "sender":
        if sender_id:
            extra["spNumber"] = sender_id
    else:
        if SMS_API_SP_NUMBER:
            extra["spNumber"] = SMS_API_SP_NUMBER

    return build_base_sms_api_payload(extra)


def get_api_status(data: Any) -> Optional[Any]:
    if not isinstance(data, dict):
        return None

    for key in (
        "status",
        "code",
        "retCode",
        "errorCode",
        "respCode",
        "resultCode",
    ):
        if key in data:
            return data.get(key)

    return None


def is_success_response(data: Any, http_status: int) -> bool:
    if http_status >= 400:
        return False

    if isinstance(data, dict):
        if data.get("success") is True:
            return True

    status = get_api_status(data)

    if status is None:
        return 200 <= http_status < 300

    if status in (0, "0", 200, "200"):
        return True

    if str(status).lower() in ("success", "ok", "true"):
        return True

    return False


async def post_sms_api(
    session: aiohttp.ClientSession,
    url: str,
    payload: Dict[str, Any],
) -> Tuple[bool, Any, str]:
    headers = build_sms_api_headers()

    try:
        async with session.post(
            url,
            json=payload,
            headers=headers,
            timeout=aiohttp.ClientTimeout(total=REQUEST_TIMEOUT),
        ) as response:
            raw_text = await response.text()

            try:
                data = await response.json(content_type=None)
            except Exception:
                data = {
                    "http_status": response.status,
                    "raw": raw_text,
                }

            ok = is_success_response(data, response.status)
            return ok, data, raw_text

    except Exception as e:
        logger.exception("Erreur requête SMS API")
        return False, {"exception": str(e)}, str(e)


async def send_sms_api_batch(
    session: aiohttp.ClientSession,
    numbers: List[str],
    sms_text: str,
    route: str,
    sender_id: Optional[str] = None,
) -> Tuple[bool, Any]:
    payload = build_payload(
        numbers=numbers,
        sms_text=sms_text,
        route=route,
        sender_id=sender_id,
    )

    ok, data, raw_text = await post_sms_api(
        session=session,
        url=SMS_API_SUBMITTAL_URL,
        payload=payload,
    )

    if ok:
        return True, data

    logger.error("Batch SMS API refusé. Réponse=%s Raw=%s", data, raw_text)
    return False, data


def sms_api_config_ok() -> bool:
    return bool(SMS_API_USERNAME and SMS_API_KEY and SMS_API_SUBMITTAL_URL)


def parse_key_value_args(raw: str) -> Dict[str, Any]:
    result: Dict[str, Any] = {}

    for part in re.split(r"\s+", raw.strip()):
        if not part or "=" not in part:
            continue
        key, value = part.split("=", 1)
        key = key.strip()
        value = value.strip()
        if key:
            result[key] = value

    return result


# ============================================================
# HANDLERS START / MENU
# ============================================================

@router.message(CommandStart())
@router.message(Command("campagne"))
async def start_campaign(message: Message, state: FSMContext):
    await state.clear()
    ensure_user(message.from_user.id)
    delete_expired_licenses()

    if not is_admin(message.from_user.id) and not has_valid_license(message.from_user.id):
        await message.answer(
            "⛔ <b>Accès refusé.</b>\n\n"
            "Tu n'as pas de licence active ou ta licence est expirée.\n"
            "Contacte @kaplanmr\n\n"
            f"Ton ID Telegram : <code>{message.from_user.id}</code>"
        )
        return

    await message.answer(
        menu_text(),
        reply_markup=route_keyboard(message.from_user.id),
    )


@router.message(Command("cancel"))
async def cancel_command(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("❌ <b>Campagne annulée.</b>")


# ============================================================
# USER COMMANDS
# ============================================================

@router.message(Command("solde", "my_balance"))
async def internal_balance_command(message: Message):
    bal = get_user_balance(message.from_user.id)

    await message.answer(
        f"💰 <b>Solde interne</b>\n\n"
        f"Disponible : <b>${money_fmt(bal)}</b>"
    )


@router.message(Command("licence", "license"))
async def license_command(message: Message):
    user_id = message.from_user.id

    if is_admin(user_id):
        await message.answer("👑 <b>Tu es admin.</b>\n\nLicence illimitée.")
        return

    expires_at = get_license_expiry(user_id)

    if not expires_at:
        await message.answer(
            "❌ <b>Aucune licence active.</b>\n\n"
            f"Ton ID Telegram : <code>{user_id}</code>"
        )
        return

    remaining = max(0, expires_at - now_ts())
    expires_dt = datetime.fromtimestamp(expires_at, tz=timezone.utc)

    await message.answer(
        "✅ <b>Licence active</b>\n\n"
        f"Expire le : <code>{expires_dt.strftime('%Y-%m-%d %H:%M:%S UTC')}</code>\n"
        f"Temps restant : <b>{remaining // 86400}</b> jours"
    )


# ============================================================
# ADMIN COMMANDS
# ============================================================

@router.message(Command("grant", "grant_license"))
async def grant_command(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔ <b>Accès refusé.</b>")
        return

    parts = message.text.split()

    if len(parts) != 3:
        await message.answer(
            "Usage :\n"
            "<code>/grant_license USER_ID DUREE</code>\n\n"
            "Exemples :\n"
            "<code>/grant_license 123456789 30d</code>\n"
            "<code>/grant_license 123456789 12h</code>\n"
            "<code>/grant_license 123456789 2w</code>\n\n"
            "Si tu mets juste <code>30</code>, ça veut dire 30 jours."
        )
        return

    try:
        user_id = int(parts[1])
        duration_raw = parts[2]
        expires_at = grant_license(user_id, duration_raw)
    except Exception as e:
        await message.answer(f"❌ Erreur : <code>{html.escape(str(e))}</code>")
        return

    expires_dt = datetime.fromtimestamp(expires_at, tz=timezone.utc)

    await message.answer(
        "✅ <b>Licence ajoutée</b>\n\n"
        f"User ID : <code>{user_id}</code>\n"
        f"Durée : <code>{html.escape(duration_raw)}</code>\n"
        f"Expire le : <code>{expires_dt.strftime('%Y-%m-%d %H:%M:%S UTC')}</code>"
    )


@router.message(Command("revoke", "revoke_license"))
async def revoke_command(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔ <b>Accès refusé.</b>")
        return

    parts = message.text.split()

    if len(parts) != 2:
        await message.answer(
            "Usage :\n"
            "<code>/revoke_license USER_ID</code>"
        )
        return

    try:
        user_id = int(parts[1])
    except ValueError:
        await message.answer("❌ USER_ID invalide.")
        return

    revoke_license(user_id)

    await message.answer(
        "✅ <b>Licence supprimée</b>\n\n"
        f"User ID : <code>{user_id}</code>"
    )


@router.message(Command("addbalance", "add_balance"))
async def add_balance_command(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔ <b>Accès refusé.</b>")
        return

    parts = message.text.split()

    if len(parts) != 3:
        await message.answer(
            "Usage :\n"
            "<code>/add_balance USER_ID MONTANT</code>\n\n"
            "Exemple :\n"
            "<code>/add_balance 123456789 10.50</code>"
        )
        return

    try:
        user_id = int(parts[1])
        amount = money(parts[2])
        new_balance = add_user_balance(user_id, amount)
    except Exception as e:
        await message.answer(f"❌ Erreur : <code>{html.escape(str(e))}</code>")
        return

    await message.answer(
        "✅ <b>Solde ajouté</b>\n\n"
        f"User ID : <code>{user_id}</code>\n"
        f"Ajout : <b>${money_fmt(amount)}</b>\n"
        f"Nouveau solde : <b>${money_fmt(new_balance)}</b>"
    )


@router.message(Command("removebalance", "remove_balance"))
async def remove_balance_command(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔ <b>Accès refusé.</b>")
        return

    parts = message.text.split()

    if len(parts) != 3:
        await message.answer(
            "Usage :\n"
            "<code>/remove_balance USER_ID MONTANT</code>\n\n"
            "Exemple :\n"
            "<code>/remove_balance 123456789 5.00</code>"
        )
        return

    try:
        user_id = int(parts[1])
        amount = money(parts[2])
        new_balance = remove_user_balance(user_id, amount)
    except Exception as e:
        await message.answer(f"❌ Erreur : <code>{html.escape(str(e))}</code>")
        return

    await message.answer(
        "✅ <b>Solde retiré</b>\n\n"
        f"User ID : <code>{user_id}</code>\n"
        f"Retrait : <b>${money_fmt(amount)}</b>\n"
        f"Nouveau solde : <b>${money_fmt(new_balance)}</b>"
    )


@router.message(Command("setbalance", "set_balance"))
async def set_balance_command(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔ <b>Accès refusé.</b>")
        return

    parts = message.text.split()

    if len(parts) != 3:
        await message.answer(
            "Usage :\n"
            "<code>/set_balance USER_ID MONTANT</code>\n\n"
            "Exemple :\n"
            "<code>/set_balance 123456789 25</code>"
        )
        return

    try:
        user_id = int(parts[1])
        amount = money(parts[2])
        new_balance = set_user_balance(user_id, amount)
    except Exception as e:
        await message.answer(f"❌ Erreur : <code>{html.escape(str(e))}</code>")
        return

    await message.answer(
        "✅ <b>Solde défini</b>\n\n"
        f"User ID : <code>{user_id}</code>\n"
        f"Nouveau solde : <b>${money_fmt(new_balance)}</b>"
    )


@router.message(Command("user_info", "userinfo"))
async def user_info_command(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔ <b>Accès refusé.</b>")
        return

    parts = message.text.split()

    if len(parts) != 2:
        await message.answer(
            "Usage :\n"
            "<code>/user_info USER_ID</code>"
        )
        return

    try:
        user_id = int(parts[1])
        info = get_user_info(user_id)
    except Exception as e:
        await message.answer(f"❌ Erreur : <code>{html.escape(str(e))}</code>")
        return

    expires_at = info["expires_at"]

    if expires_at:
        expires_dt = datetime.fromtimestamp(expires_at, tz=timezone.utc)
        license_text = expires_dt.strftime("%Y-%m-%d %H:%M:%S UTC")
    else:
        license_text = "Aucune"

    await message.answer(
        "👤 <b>User info</b>\n\n"
        f"User ID : <code>{info['user_id']}</code>\n"
        f"Admin : <b>{'Oui' if info['is_admin'] else 'Non'}</b>\n"
        f"Licence active : <b>{'Oui' if info['has_license'] else 'Non'}</b>\n"
        f"Expiration licence : <code>{license_text}</code>\n"
        f"Solde : <b>${money_fmt(info['balance'])}</b>\n\n"
        f"SMS total ledger : <b>{info['total_sms']}</b>\n"
        f"SMS accepted/pending : <b>{info['accepted_sms']}</b>\n"
        f"SMS delivered : <b>{info['delivered_sms']}</b>\n"
        f"SMS refunded : <b>{info['refunded_sms']}</b>"
    )


# ============================================================
# API COMMANDS
# ============================================================

@router.message(Command("balance"))
async def balance_command(message: Message):
    if not SMS_API_USERNAME or not SMS_API_KEY or not SMS_API_BALANCE_URL:
        await message.answer(
            "❌ <b>Configuration SMS API manquante.</b>\n\n"
            "Vérifie :\n"
            "<code>SMS_API_DOMAIN</code>\n"
            "<code>SMS_API_USERNAME</code>\n"
            "<code>SMS_API_KEY</code>"
        )
        return

    payload = build_base_sms_api_payload()

    async with aiohttp.ClientSession() as session:
        ok, data, raw_text = await post_sms_api(session, SMS_API_BALANCE_URL, payload)

    if ok:
        text = json_pretty(data)
        await message.answer(
            "💰 <b>Balance API</b>\n\n"
            f"<code>{html.escape(truncate_text(text))}</code>"
        )
    else:
        await message.answer(
            "❌ <b>Erreur balance API</b>\n\n"
            f"<code>{html.escape(truncate_text(raw_text or json_pretty(data)))}</code>"
        )


@router.message(Command("records"))
async def records_command(message: Message):
    if not SMS_API_USERNAME or not SMS_API_KEY or not SMS_API_RECORDS_URL:
        await message.answer(
            "❌ <b>Configuration SMS API manquante.</b>\n\n"
            "Vérifie :\n"
            "<code>SMS_API_DOMAIN</code>\n"
            "<code>SMS_API_USERNAME</code>\n"
            "<code>SMS_API_KEY</code>"
        )
        return

    raw_args = ""
    if message.text:
        parts = message.text.split(maxsplit=1)
        if len(parts) > 1:
            raw_args = parts[1]

    extra = parse_key_value_args(raw_args)
    payload = build_base_sms_api_payload(extra)

    async with aiohttp.ClientSession() as session:
        ok, data, raw_text = await post_sms_api(session, SMS_API_RECORDS_URL, payload)

    if ok:
        settlement = auto_settle_from_records(data)
        text = json_pretty(data)

        await message.answer(
            "📑 <b>Records API</b>\n\n"
            f"🔎 Statuts détectés : <b>{settlement['found']}</b>\n"
            f"✅ Delivered confirmés : <b>{settlement['delivered']}</b>\n"
            f"↩️ Refunded failed : <b>{settlement['refunded']}</b>\n"
            f"⏳ Pending/inconnus : <b>{settlement['pending']}</b>\n\n"
            f"<code>{html.escape(truncate_text(text))}</code>"
        )
    else:
        await message.answer(
            "❌ <b>Erreur records API</b>\n\n"
            f"<code>{html.escape(truncate_text(raw_text or json_pretty(data)))}</code>"
        )


# ============================================================
# CALLBACKS MENU
# ============================================================

@router.callback_query(F.data == "user:balance")
async def user_balance_callback(callback: CallbackQuery):
    bal = get_user_balance(callback.from_user.id)

    await callback.answer()

    if callback.message:
        await callback.message.answer(
            f"💰 <b>Ton solde interne</b>\n\n"
            f"Solde disponible : <b>${money_fmt(bal)}</b>"
        )


@router.callback_query(F.data == "user:license")
async def user_license_callback(callback: CallbackQuery):
    user_id = callback.from_user.id

    await callback.answer()

    if is_admin(user_id):
        text = "👑 <b>Tu es admin.</b>\n\nLicence illimitée."
    else:
        expires_at = get_license_expiry(user_id)

        if not expires_at:
            text = "❌ <b>Aucune licence active.</b>"
        else:
            remaining = max(0, expires_at - now_ts())
            expires_dt = datetime.fromtimestamp(expires_at, tz=timezone.utc)

            text = (
                "✅ <b>Licence active</b>\n\n"
                f"Expire le : <code>{expires_dt.strftime('%Y-%m-%d %H:%M:%S UTC')}</code>\n"
                f"Temps restant : <b>{remaining // 86400}</b> jours"
            )

    if callback.message:
        await callback.message.answer(text)


@router.callback_query(F.data == "admin:panel")
async def admin_panel_callback(callback: CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("Accès refusé.", show_alert=True)
        return

    await callback.answer()

    if callback.message:
        await callback.message.answer(
            "👑 <b>Panel admin</b>\n\n"
            "Commandes disponibles :\n\n"
            "<code>/grant_license USER_ID DUREE</code>\n"
            "Exemple : <code>/grant_license 123456789 30d</code>\n\n"
            "<code>/revoke_license USER_ID</code>\n"
            "Supprime une licence.\n\n"
            "<code>/add_balance USER_ID MONTANT</code>\n"
            "Exemple : <code>/add_balance 123456789 10.50</code>\n\n"
            "<code>/remove_balance USER_ID MONTANT</code>\n"
            "Exemple : <code>/remove_balance 123456789 5</code>\n\n"
            "<code>/set_balance USER_ID MONTANT</code>\n"
            "Exemple : <code>/set_balance 123456789 25</code>\n\n"
            "<code>/user_info USER_ID</code>\n"
            "Voir infos user.\n\n"
            "<code>/balance</code>\n"
            "Voir la balance API SMS.\n\n"
            "<code>/records</code>\n"
            "Consulter records API + sync refunds."
        )


@router.callback_query(F.data == "campaign:cancel")
async def cancel_callback(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    await callback.answer("Campagne annulée.")
    if callback.message:
        await callback.message.edit_text("❌ <b>Campagne annulée.</b>")


@router.callback_query(F.data == "campaign:restart")
async def restart_callback(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    await callback.answer()
    if callback.message:
        await callback.message.edit_text(
            menu_text(),
            reply_markup=route_keyboard(callback.from_user.id),
        )


# ============================================================
# ROUTE SELECTION
# ============================================================

@router.callback_query(F.data == "route:sim")
async def choose_route_sim(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    await state.update_data(route="sim", sender_id=None)
    await state.set_state(CampaignState.waiting_numbers)

    await callback.answer("Route SIM sélectionnée 📲")

    if callback.message:
        await callback.message.edit_text(
            "📲 <b>Route SIM sélectionnée</b>\n\n" + numbers_help_text(),
            reply_markup=cancel_keyboard(),
        )


@router.callback_query(F.data == "route:sender")
async def choose_route_sender(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    await state.update_data(route="sender")
    await state.set_state(CampaignState.waiting_sender_id)

    await callback.answer("Route Sender ID sélectionnée 🏷️")

    if callback.message:
        await callback.message.edit_text(
            "🏷️ <b>Route Sender ID sélectionnée</b>\n\n"
            "Envoie maintenant le <b>Sender ID</b> à utiliser.\n\n"
            "Maximum : <b>20 caractères</b>.",
            reply_markup=cancel_keyboard(),
        )


# ============================================================
# SENDER ID
# ============================================================

@router.message(CampaignState.waiting_sender_id)
async def receive_sender_id(message: Message, state: FSMContext):
    sender_id = message.text.strip() if message.text else ""

    if not sender_id:
        await message.answer("⚠️ Sender ID vide. Envoie un Sender ID valide.")
        return

    if len(sender_id) > 20:
        await message.answer(
            "⚠️ Le Sender ID est trop long.\n"
            "Maximum : <b>20 caractères</b>."
        )
        return

    if "\n" in sender_id or "\r" in sender_id:
        await message.answer("⚠️ Le Sender ID doit être sur une seule ligne.")
        return

    await state.update_data(sender_id=sender_id)
    await state.set_state(CampaignState.waiting_numbers)

    await message.answer(
        f"✅ Sender ID enregistré : <code>{html.escape(sender_id)}</code>\n\n"
        + numbers_help_text(),
        reply_markup=cancel_keyboard(),
    )


# ============================================================
# NUMBERS
# ============================================================

@router.message(CampaignState.waiting_numbers)
async def receive_numbers(message: Message, state: FSMContext):
    input_text, source_label, error = await extract_numbers_input(message)

    if error:
        await message.answer(error, reply_markup=cancel_keyboard())
        return

    if not input_text:
        await message.answer(
            "⚠️ Aucun contenu détecté.\n\n"
            "Envoie les numéros en message direct, fichier <code>.txt</code> ou fichier <code>.xlsx</code>.",
            reply_markup=cancel_keyboard(),
        )
        return

    numbers, duplicates, invalid = parse_numbers(input_text)

    if not numbers:
        await message.answer(
            "⚠️ Aucun numéro valide détecté.\n\n"
            "Exemple :\n"
            "<code>33612345678\n33712345678\n33812345678</code>\n\n"
            "Tu peux aussi envoyer un fichier <code>.txt</code> ou <code>.xlsx</code>.",
            reply_markup=cancel_keyboard(),
        )
        return

    if len(numbers) > MAX_NUMBERS:
        await message.answer(
            f"⚠️ Trop de numéros.\n\n"
            f"Maximum autorisé : <b>{MAX_NUMBERS}</b>\n"
            f"Reçus valides : <b>{len(numbers)}</b>\n\n"
            "Réduis la liste puis renvoie-la.",
            reply_markup=cancel_keyboard(),
        )
        return

    await state.update_data(numbers=numbers)
    await state.set_state(CampaignState.waiting_text)

    preview = "\n".join(numbers[:10])
    more = ""
    if len(numbers) > 10:
        more = f"\n... +{len(numbers) - 10} autres"

    await message.answer(
        "✅ <b>Numéros enregistrés</b>\n\n"
        f"📥 Source : <b>{source_label}</b>\n"
        f"📞 Valides : <b>{len(numbers)}</b>\n"
        f"♻️ Doublons supprimés : <b>{duplicates}</b>\n"
        f"⚠️ Invalides ignorés : <b>{invalid}</b>\n\n"
        f"<b>Aperçu :</b>\n<code>{html.escape(preview)}</code>{more}\n\n"
        "✉️ Envoie maintenant le texte du SMS.",
        reply_markup=cancel_keyboard(),
    )


# ============================================================
# SMS TEXT
# ============================================================

@router.message(CampaignState.waiting_text)
async def receive_sms_text(message: Message, state: FSMContext):
    sms_text = message.text or ""
    sms_text = sms_text.strip()

    if not sms_text:
        await message.answer("⚠️ Le texte du SMS est vide. Envoie un texte valide.")
        return

    await state.update_data(sms_text=sms_text)
    await state.set_state(CampaignState.waiting_confirmation)

    data = await state.get_data()

    route = data.get("route")
    sender_id = data.get("sender_id")
    numbers = data.get("numbers", [])

    if route == "sim":
        route_label = "📲 Route SIM"
        sender_label = SMS_API_SP_NUMBER or "Non utilisé"
    else:
        route_label = "🏷️ Sender ID"
        sender_label = sender_id or ""

    price_per_sms = get_sms_price(route)
    total_cost = get_sms_total_cost(route, len(numbers))
    current_balance = get_user_balance(message.from_user.id)

    await message.answer(
        "📋 <b>Confirmation campagne SMS</b>\n\n"
        f"🚦 Route : <b>{route_label}</b>\n"
        f"🏷️ Sender / SP Number : <code>{html.escape(sender_label)}</code>\n"
        f"📞 Nombre de numéros : <b>{len(numbers)}</b>\n"
        f"📦 Batch size : <b>{BATCH_SIZE}</b>\n"
        f"💵 Prix / SMS : <b>${money_fmt(price_per_sms)}</b>\n"
        f"💸 Coût total : <b>${money_fmt(total_cost)}</b>\n"
        f"💰 Ton solde : <b>${money_fmt(current_balance)}</b>\n\n"
        "✉️ <b>Message :</b>\n"
        f"<code>{html.escape(sms_text)}</code>\n\n"
        "Confirme l’envoi :",
        reply_markup=confirmation_keyboard(),
    )


@router.callback_query(F.data == "campaign:edit_text")
async def edit_text_callback(callback: CallbackQuery, state: FSMContext):
    current_state = await state.get_state()

    if current_state != CampaignState.waiting_confirmation.state:
        await callback.answer("Action impossible maintenant.", show_alert=True)
        return

    await state.set_state(CampaignState.waiting_text)
    await callback.answer()

    if callback.message:
        await callback.message.edit_text(
            "✏️ <b>Modification du texte SMS</b>\n\n"
            "Envoie le nouveau texte du SMS.",
            reply_markup=cancel_keyboard(),
        )


# ============================================================
# SEND CAMPAIGN
# ============================================================

@router.callback_query(F.data == "campaign:send")
async def send_campaign(callback: CallbackQuery, state: FSMContext):
    current_state = await state.get_state()

    if current_state != CampaignState.waiting_confirmation.state:
        await callback.answer("Aucune campagne prête à envoyer.", show_alert=True)
        return

    user_id = callback.from_user.id

    data = await state.get_data()

    route = data.get("route")
    sender_id = data.get("sender_id")
    numbers: List[str] = data.get("numbers", [])
    sms_text = data.get("sms_text", "")

    if not route or not numbers or not sms_text:
        await callback.answer("Données campagne incomplètes.", show_alert=True)
        return

    if not sms_api_config_ok():
        await callback.answer("Config SMS API manquante.", show_alert=True)
        if callback.message:
            await callback.message.answer(
                "❌ <b>Configuration SMS API manquante.</b>\n\n"
                "Vérifie dans ton fichier <code>.env</code> :\n"
                "<code>SMS_API_DOMAIN</code>\n"
                "<code>SMS_API_USERNAME</code>\n"
                "<code>SMS_API_KEY</code>"
            )
        return

    total_cost = get_sms_total_cost(route, len(numbers))
    current_balance = get_user_balance(user_id)

    if current_balance < total_cost:
        await callback.answer("Solde insuffisant.", show_alert=True)
        if callback.message:
            await callback.message.answer(
                "❌ <b>Solde insuffisant</b>\n\n"
                f"SMS : <b>{len(numbers)}</b>\n"
                f"Coût total : <b>${money_fmt(total_cost)}</b>\n"
                f"Ton solde : <b>${money_fmt(current_balance)}</b>\n\n"
                "Recharge ton solde avant d'envoyer."
            )
        return

    await callback.answer("Envoi lancé 🚀")

    campaign_id = uuid.uuid4().hex

    batches = chunk_list(numbers, BATCH_SIZE)
    total_batches = len(batches)

    accepted_total = 0
    refused_total = 0
    refunded_total = 0
    charged_total = money("0")

    route_label = "📲 SIM" if route == "sim" else f"🏷️ Sender ID: {html.escape(sender_id or '')}"

    if callback.message:
        progress_msg = await callback.message.edit_text(
            "🚀 <b>Campagne SMS lancée</b>\n\n"
            f"🆔 Campaign ID : <code>{campaign_id}</code>\n"
            f"🚦 Route : <b>{route_label}</b>\n"
            f"📞 Total : <b>{len(numbers)}</b>\n"
            f"📦 Batches : <b>{total_batches}</b>\n"
            f"💸 Coût max : <b>${money_fmt(total_cost)}</b>\n\n"
            "⏳ Préparation de l’envoi..."
        )
    else:
        progress_msg = None

    async with aiohttp.ClientSession() as session:
        for index, batch in enumerate(batches, start=1):
            try:
                ledger_ids, batch_cost = reserve_sms_balance(
                    user_id=user_id,
                    route=route,
                    phones=batch,
                    campaign_id=campaign_id,
                )
                charged_total = money(charged_total + batch_cost)

            except ValueError as e:
                refused_total += len(batch)
                logger.error("Réservation solde impossible: %s", e)

                if progress_msg:
                    try:
                        await progress_msg.edit_text(
                            "❌ <b>Campagne arrêtée</b>\n\n"
                            f"Erreur : <code>{html.escape(str(e))}</code>"
                        )
                    except Exception:
                        pass
                break

            ok, response_data = await send_sms_api_batch(
                session=session,
                numbers=batch,
                sms_text=sms_text,
                route=route,
                sender_id=sender_id,
            )

            if ok:
                accepted_total += len(batch)
                mark_campaign_batch_accepted(ledger_ids)
                batch_status = "✅ Accepté API"
            else:
                refused_total += len(batch)
                refunded_count, refunded_amount = refund_ledger_ids(
                    ledger_ids,
                    provider_status="api_refused",
                )
                refunded_total += refunded_count
                charged_total = money(charged_total - refunded_amount)
                batch_status = "❌ Refusé API / refund"

            logger.info(
                "Campaign=%s | Batch %s/%s | size=%s | ok=%s | response=%s",
                campaign_id,
                index,
                total_batches,
                len(batch),
                ok,
                response_data,
            )

            current_balance_after = get_user_balance(user_id)

            progress_text = (
                "🚀 <b>Campagne SMS en cours</b>\n\n"
                f"🆔 Campaign ID : <code>{campaign_id}</code>\n"
                f"🚦 Route : <b>{route_label}</b>\n"
                f"📦 Batch : <b>{index}/{total_batches}</b>\n"
                f"📨 Dernier batch : <b>{batch_status}</b>\n\n"
                f"✅ Acceptés API : <b>{accepted_total}</b>\n"
                f"❌ Refusés API : <b>{refused_total}</b>\n"
                f"↩️ Refund immédiat : <b>{refunded_total}</b>\n"
                f"💸 Débité net : <b>${money_fmt(charged_total)}</b>\n"
                f"💰 Solde restant : <b>${money_fmt(current_balance_after)}</b>\n"
                f"📞 Total : <b>{len(numbers)}</b>"
            )

            if progress_msg:
                try:
                    await progress_msg.edit_text(progress_text)
                except Exception:
                    pass

            await asyncio.sleep(0.2)

    await state.clear()

    final_balance = get_user_balance(user_id)

    final_text = (
        "✅ <b>Campagne terminée</b>\n\n"
        f"🆔 Campaign ID : <code>{campaign_id}</code>\n"
        f"🚦 Route : <b>{route_label}</b>\n"
        f"📞 Total : <b>{len(numbers)}</b>\n"
        f"✅ Acceptés API : <b>{accepted_total}</b>\n"
        f"❌ Refusés API : <b>{refused_total}</b>\n"
        f"↩️ Refund immédiat : <b>{refunded_total}</b>\n"
        f"💸 Débité net : <b>${money_fmt(charged_total)}</b>\n"
        f"💰 Solde restant : <b>${money_fmt(final_balance)}</b>\n\n"
        "ℹ️ Si un SMS accepté API revient plus tard en failed/undelivered dans "
        "<code>/records</code>, il sera refund automatiquement."
    )

    if progress_msg:
        try:
            await progress_msg.edit_text(final_text)
        except Exception:
            if callback.message:
                await callback.message.answer(final_text)
    elif callback.message:
        await callback.message.answer(final_text)


# ============================================================
# FALLBACK
# ============================================================

@router.message()
async def fallback(message: Message):
    if is_admin(message.from_user.id):
        admin_text = (
            "\n\n👑 Commandes admin :\n"
            "/grant_license USER_ID DUREE - Ajouter une licence\n"
            "/revoke_license USER_ID - Supprimer une licence\n"
            "/add_balance USER_ID MONTANT - Ajouter du solde\n"
            "/remove_balance USER_ID MONTANT - Retirer du solde\n"
            "/set_balance USER_ID MONTANT - Définir le solde\n"
            "/user_info USER_ID - Infos utilisateur"
        )
    else:
        admin_text = ""

    await message.answer(
        "📨 Utilise <b>/start</b> ou <b>/campagne</b> pour lancer une campagne SMS.\n\n"
        "Commandes disponibles :\n"
        "/my_balance - Voir ton solde interne\n"
        "/licence - Voir ta licence\n"
        "/balance - Voir le solde API\n"
        "/records - Consulter les records API"
        f"{admin_text}"
    )


# ============================================================
# MAIN
# ============================================================

async def main():
    init_app_db()
    delete_expired_licenses()
    asyncio.create_task(license_cleanup_loop())

    logger.info("Bot démarré.")

    if not ADMIN_IDS:
        logger.warning("ADMIN_IDS est vide. Aucun admin configuré.")
    else:
        logger.info("Admins configurés : %s", sorted(ADMIN_IDS))

    logger.info("DB_PATH=%s", DB_PATH)
    logger.info("SMS_API_SUBMITTAL_URL=%s", SMS_API_SUBMITTAL_URL)
    logger.info("SMS_API_RECORDS_URL=%s", SMS_API_RECORDS_URL)
    logger.info("SMS_API_BALANCE_URL=%s", SMS_API_BALANCE_URL)
    logger.info("BATCH_SIZE=%s MAX_NUMBERS=%s", BATCH_SIZE, MAX_NUMBERS)
    logger.info("SMS_API_KEEP_PLUS=%s", SMS_API_KEEP_PLUS)
    logger.info("SMS_API_TIMESTAMP_UNIT=%s", SMS_API_TIMESTAMP_UNIT)
    logger.info("SMS_API_SIGN_TYPE=%s", SMS_API_SIGN_TYPE)
    logger.info("SMS_API_SIGN_UPPER=%s", SMS_API_SIGN_UPPER)
    logger.info("MAX_UPLOAD_MB=%s", MAX_UPLOAD_MB)
    logger.info("SMS_PRICE_SIM=%s SMS_PRICE_SENDER=%s", SMS_PRICE_SIM, SMS_PRICE_SENDER)

    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
